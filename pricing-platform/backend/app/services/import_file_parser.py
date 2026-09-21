import csv
import io
from collections.abc import Iterator

import openpyxl

from app.models.import_job import ImportFileType


def detect_file_type(filename: str) -> ImportFileType | None:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return ImportFileType.CSV
    if lower.endswith(".xlsx"):
        return ImportFileType.XLSX
    return None


def parse_header(content: bytes, file_type: ImportFileType) -> list[str]:
    """Returns the file's header row — Step 2's "preview detected columns"."""
    if file_type == ImportFileType.CSV:
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        return [c.strip() for c in next(reader, [])]

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        row = next(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True), ())
    finally:
        workbook.close()
    return [str(cell).strip() if cell is not None else "" for cell in row]


def iter_rows(content: bytes, file_type: ImportFileType) -> Iterator[dict[str, str]]:
    """Yields one dict per data row, keyed by the header's column names,
    with every value coerced to a (possibly empty) string for uniform
    downstream parsing regardless of source format. Blank rows are skipped.
    """
    if file_type == ImportFileType.CSV:
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        for row in reader:
            if any((v or "").strip() for v in row.values()):
                yield {(k or "").strip(): (v or "").strip() for k, v in row.items()}
        return

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows_iter = workbook.active.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows_iter, ())]
        for row in rows_iter:
            if all(cell is None for cell in row):
                continue
            values = [("" if cell is None else str(cell).strip()) for cell in row]
            yield dict(zip(header, values, strict=False))
    finally:
        workbook.close()
