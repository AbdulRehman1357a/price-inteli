import io
import logging
import socket
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.exceptions import ConflictError, NotFoundError, ValidationError
from app.core.security import hash_password
from app.db.session import SessionLocal
from app.models.import_job import ImportEntityType, ImportFileType, ImportJob, ImportStatus
from app.models.import_row_error import ImportRowError
from app.models.inventory import Inventory
from app.models.product import Product
from app.models.user import User, UserStatus
from app.repositories.category_repository import CategoryRepository
from app.repositories.import_job_repository import ImportJobRepository
from app.repositories.import_row_error_repository import ImportRowErrorRepository
from app.repositories.inventory_repository import InventoryRepository
from app.repositories.product_repository import ProductRepository
from app.repositories.role_repository import RoleRepository
from app.repositories.store_repository import StoreRepository
from app.repositories.user_repository import UserRepository
from app.repositories.user_role_repository import UserRoleRepository
from app.schemas.import_job import ImportField, ImportMappingRequest, ImportRowErrorListParams
from app.services.import_file_parser import detect_file_type, iter_rows, parse_header
from app.storage import get_storage_adapter

logger = logging.getLogger(__name__)

_COMMIT_EVERY = 100  # rows between progress commits, so polling GET /imports/{id} sees live progress

_REQUIRED_FIELDS: dict[ImportEntityType, list[ImportField]] = {
    ImportEntityType.PRODUCTS: [
        ImportField.SKU,
        ImportField.PRODUCT_NAME,
        ImportField.CATEGORY,
        ImportField.SELLING_PRICE,
    ],
    ImportEntityType.INVENTORY: [ImportField.SKU, ImportField.QUANTITY],
    ImportEntityType.USERS: [
        ImportField.FIRST_NAME,
        ImportField.LAST_NAME,
        ImportField.EMAIL,
        ImportField.PASSWORD,
        ImportField.ROLE,
    ],
}

# (column label, example value) per entity_type — the label doubles as the
# downloadable template's header row, and the example row shows the
# expected format (e.g. decimal places, "store" is a code not a name).
_TEMPLATE_COLUMNS: dict[ImportEntityType, list[tuple[str, str]]] = {
    ImportEntityType.PRODUCTS: [
        ("SKU", "SAMPLE-SKU-001"),
        ("Barcode", "012345678905"),
        ("Product Name", "Sample Product"),
        ("Category", "Electronics"),
        ("Cost Price", "10.0000"),
        ("Selling Price", "19.9900"),
    ],
    ImportEntityType.INVENTORY: [
        ("SKU", "SAMPLE-SKU-001"),
        ("Store", "STORE-CODE"),
        ("Quantity", "100"),
        ("Reorder Point", "20"),
    ],
    ImportEntityType.USERS: [
        ("First Name", "Jane"),
        ("Last Name", "Doe"),
        ("Email", "jane.doe@example.com"),
        ("Password", "ChangeMe123!"),
        ("Role", "Store Manager"),
        ("Phone", "+1-555-0100"),
    ],
}


def build_import_template(entity_type: ImportEntityType) -> tuple[str, bytes]:
    """Builds a downloadable .xlsx template: a bold header row matching the
    exact labels shown in the Map Columns step, plus one highlighted
    example row to illustrate the expected format. Returns (filename, bytes).
    """
    columns = _TEMPLATE_COLUMNS[entity_type]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = entity_type.value.capitalize()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    example_font = Font(italic=True, color="806000")
    example_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    note_font = Font(italic=True, color="999999", size=9)

    for col_index, (label, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=label)
        cell.font = header_font
        cell.fill = header_fill
        sheet.column_dimensions[cell.column_letter].width = max(len(label) + 6, 16)

    for col_index, (_, sample) in enumerate(columns, start=1):
        cell = sheet.cell(row=2, column=col_index, value=sample)
        cell.font = example_font
        cell.fill = example_fill

    sheet.cell(
        row=2,
        column=len(columns) + 2,
        value="<- example row: replace with real data or delete before uploading",
    ).font = note_font
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = f"{entity_type.value}_import_template.xlsx"
    return filename, buffer.getvalue()


def create_import_job(
    db: Session,
    *,
    organization_id: uuid.UUID,
    user_id: uuid.UUID,
    filename: str,
    content: bytes,
    entity_type: ImportEntityType,
    store_id: uuid.UUID | None,
) -> ImportJob:
    file_type = detect_file_type(filename)
    if file_type is None:
        raise ValidationError("Only .csv and .xlsx files are supported.")

    if store_id is not None:
        if StoreRepository(db).get_by_id_for_organization(store_id, organization_id) is None:
            raise NotFoundError("Store not found.", code="store_not_found")

    try:
        header = parse_header(content, file_type)
    except Exception as exc:  # noqa: BLE001 — any parse failure is a client-facing validation error
        raise ValidationError(f"Could not read this file: {exc}") from exc
    if not header:
        raise ValidationError("This file has no header row.")

    key = f"imports/{organization_id}/{uuid.uuid4()}_{filename}"
    content_type = (
        "text/csv"
        if file_type == ImportFileType.CSV
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    file_url = get_storage_adapter().upload(key=key, content=content, content_type=content_type)

    job = ImportJob(
        id=uuid.uuid4(),
        organization_id=organization_id,
        store_id=store_id,
        file_name=filename,
        file_type=file_type,
        file_url=file_url,
        entity_type=entity_type,
        status=ImportStatus.UPLOADED,
        created_by=user_id,
        detected_columns=header,
    )
    return ImportJobRepository(db).add(job)


def get_import_job(db: Session, *, organization_id: uuid.UUID, import_job_id: uuid.UUID) -> ImportJob:
    job = ImportJobRepository(db).get_by_id_for_organization(import_job_id, organization_id)
    if job is None:
        raise NotFoundError("Import job not found.", code="import_job_not_found")
    return job


def list_row_errors(
    db: Session, *, organization_id: uuid.UUID, import_job_id: uuid.UUID, params: ImportRowErrorListParams
) -> tuple[list[ImportRowError], int]:
    get_import_job(db, organization_id=organization_id, import_job_id=import_job_id)  # tenant check
    offset = (params.page - 1) * params.page_size
    return ImportRowErrorRepository(db).list_for_job(import_job_id, offset=offset, limit=params.page_size)


def start_import(
    db: Session, *, organization_id: uuid.UUID, import_job_id: uuid.UUID, payload: ImportMappingRequest
) -> ImportJob:
    job = get_import_job(db, organization_id=organization_id, import_job_id=import_job_id)
    if job.status != ImportStatus.UPLOADED:
        raise ConflictError(
            f"This import has already been started (status: {job.status}).", code="import_already_started"
        )

    mapping = {field.value: column for field, column in payload.mapping.items()}
    detected = set(job.detected_columns or [])
    for column in mapping.values():
        if column not in detected:
            raise ValidationError(f"'{column}' is not a column in the uploaded file.")

    required = _REQUIRED_FIELDS[job.entity_type]
    missing = [f.value for f in required if f.value not in mapping]
    needs_store_column = job.entity_type == ImportEntityType.INVENTORY and job.store_id is None
    if needs_store_column and ImportField.STORE.value not in mapping:
        missing.append(ImportField.STORE.value)
    if missing:
        raise ValidationError(f"Missing required mapping(s): {', '.join(missing)}.")

    job.column_mapping = mapping
    job.status = ImportStatus.QUEUED
    ImportJobRepository(db).add(job)

    if _broker_reachable():
        try:
            from app.tasks.imports import process_import_job_task

            process_import_job_task.delay(str(job.id))
        except Exception:  # noqa: BLE001 — broker went away between the check and the dispatch
            logger.warning("Could not dispatch import job %s to Celery.", job.id)
    else:
        logger.warning(
            "Redis broker unreachable — import job %s will stay queued until a worker is available.",
            job.id,
        )

    return job


def _broker_reachable(timeout: float = 0.5) -> bool:
    """A quick TCP pre-check before calling Celery's .delay(): Kombu's redis
    transport doesn't reliably honor broker_connection_timeout for publish
    calls, so with no broker running .delay() can block far longer than is
    acceptable inside an HTTP request — this keeps that path fast and
    predictable instead.
    """
    url = urlparse(get_settings().redis_url)
    try:
        with socket.create_connection((url.hostname or "localhost", url.port or 6379), timeout=timeout):
            return True
    except OSError:
        return False


@dataclass
class _RowValidationError(Exception):
    field_name: str | None
    message: str


def _extract(mapping: dict[str, str], row: dict[str, str], field: ImportField) -> str:
    column = mapping.get(field.value)
    if column is None:
        return ""
    return (row.get(column) or "").strip()


def _parse_decimal(value: str, *, field_name: str) -> Decimal:
    try:
        parsed = Decimal(value)
    except InvalidOperation as exc:
        raise _RowValidationError(field_name, f"'{value}' is not a valid number.") from exc
    if parsed < 0:
        raise _RowValidationError(field_name, f"{field_name} cannot be negative.")
    return parsed


def _process_products_row(
    db: Session,
    *,
    organization_id: uuid.UUID,
    mapping: dict[str, str],
    row: dict[str, str],
    seen_skus: set[str],
) -> None:
    sku = _extract(mapping, row, ImportField.SKU)
    if not sku:
        raise _RowValidationError("sku", "SKU is required.")
    if sku in seen_skus:
        raise _RowValidationError("sku", f"Duplicate SKU '{sku}' in this file.")
    seen_skus.add(sku)

    product_name = _extract(mapping, row, ImportField.PRODUCT_NAME)
    if not product_name:
        raise _RowValidationError("product_name", "Product Name is required.")

    category_name = _extract(mapping, row, ImportField.CATEGORY)
    if not category_name:
        raise _RowValidationError("category", "Category is required.")
    category = CategoryRepository(db).get_by_name(organization_id, category_name)
    if category is None:
        raise _RowValidationError("category", f"Unknown category '{category_name}'.")

    selling_price_raw = _extract(mapping, row, ImportField.SELLING_PRICE)
    if not selling_price_raw:
        raise _RowValidationError("selling_price", "Selling Price is required.")
    selling_price = _parse_decimal(selling_price_raw, field_name="selling_price")

    barcode = _extract(mapping, row, ImportField.BARCODE) or None
    cost_price_raw = _extract(mapping, row, ImportField.COST_PRICE)
    cost_price = _parse_decimal(cost_price_raw, field_name="cost_price") if cost_price_raw else None

    product_repo = ProductRepository(db)
    with db.begin_nested():
        existing = product_repo.get_by_sku(organization_id, sku)
        if existing is not None:
            existing.product_name = product_name
            existing.category_id = category.id
            existing.selling_price = selling_price
            existing.barcode = barcode
            existing.cost_price = cost_price
            product_repo.add(existing)
        else:
            product_repo.add(
                Product(
                    id=uuid.uuid4(),
                    organization_id=organization_id,
                    category_id=category.id,
                    sku=sku,
                    barcode=barcode,
                    product_name=product_name,
                    cost_price=cost_price,
                    selling_price=selling_price,
                )
            )


def _process_inventory_row(
    db: Session,
    *,
    organization_id: uuid.UUID,
    job: ImportJob,
    mapping: dict[str, str],
    row: dict[str, str],
    seen_keys: set[tuple[str, str]],
) -> None:
    sku = _extract(mapping, row, ImportField.SKU)
    if not sku:
        raise _RowValidationError("sku", "SKU is required.")
    product = ProductRepository(db).get_by_sku(organization_id, sku)
    if product is None:
        raise _RowValidationError("sku", f"Unknown SKU '{sku}' — the product must exist first.")

    if job.store_id is not None:
        store_id = job.store_id
        store_label = str(store_id)
    else:
        store_code = _extract(mapping, row, ImportField.STORE)
        if not store_code:
            raise _RowValidationError("store", "Store is required.")
        store = StoreRepository(db).get_by_code(organization_id, store_code)
        if store is None:
            raise _RowValidationError("store", f"Unknown store '{store_code}'.")
        store_id = store.id
        store_label = store_code

    dedupe_key = (sku, store_label)
    if dedupe_key in seen_keys:
        raise _RowValidationError("sku", f"Duplicate SKU '{sku}' for store '{store_label}' in this file.")
    seen_keys.add(dedupe_key)

    quantity_raw = _extract(mapping, row, ImportField.QUANTITY)
    if not quantity_raw:
        raise _RowValidationError("quantity", "Quantity is required.")
    quantity = _parse_decimal(quantity_raw, field_name="quantity")

    reorder_point_raw = _extract(mapping, row, ImportField.REORDER_POINT)
    reorder_point = (
        _parse_decimal(reorder_point_raw, field_name="reorder_point") if reorder_point_raw else None
    )

    inventory_repo = InventoryRepository(db)
    with db.begin_nested():
        existing = inventory_repo.get_by_store_and_product(organization_id, store_id, product.id)
        if existing is not None:
            existing.quantity_on_hand = quantity
            if reorder_point is not None:
                existing.reorder_point = reorder_point
            existing.quantity_available = existing.quantity_on_hand - existing.quantity_reserved
            existing.last_stock_update_at = datetime.now(UTC)
            inventory_repo.add(existing)
        else:
            inventory_repo.add(
                Inventory(
                    id=uuid.uuid4(),
                    organization_id=organization_id,
                    store_id=store_id,
                    product_id=product.id,
                    quantity_on_hand=quantity,
                    quantity_available=quantity,
                    reorder_point=reorder_point,
                    last_stock_update_at=datetime.now(UTC),
                )
            )


def _process_users_row(
    db: Session,
    *,
    organization_id: uuid.UUID,
    mapping: dict[str, str],
    row: dict[str, str],
    seen_emails: set[str],
) -> None:
    first_name = _extract(mapping, row, ImportField.FIRST_NAME)
    if not first_name:
        raise _RowValidationError("first_name", "First Name is required.")

    last_name = _extract(mapping, row, ImportField.LAST_NAME)
    if not last_name:
        raise _RowValidationError("last_name", "Last Name is required.")

    email = _extract(mapping, row, ImportField.EMAIL).lower()
    if not email:
        raise _RowValidationError("email", "Email is required.")
    if email in seen_emails:
        raise _RowValidationError("email", f"Duplicate email '{email}' in this file.")
    seen_emails.add(email)
    # Email must be globally unique, not just per-organization — see the
    # same rule enforced by user_service.create_user.
    if UserRepository(db).get_by_email(email) is not None:
        raise _RowValidationError("email", f"A user with email '{email}' already exists.")

    password = _extract(mapping, row, ImportField.PASSWORD)
    if len(password) < 8:
        raise _RowValidationError("password", "Password must be at least 8 characters.")

    role_name = _extract(mapping, row, ImportField.ROLE)
    if not role_name:
        raise _RowValidationError("role", "Role is required.")
    assignable_roles = RoleRepository(db).list_assignable(organization_id)
    role = next((r for r in assignable_roles if r.name.lower() == role_name.lower()), None)
    if role is None:
        raise _RowValidationError("role", f"Unknown role '{role_name}'.")

    phone = _extract(mapping, row, ImportField.PHONE) or None

    user_repo = UserRepository(db)
    with db.begin_nested():
        user = user_repo.add(
            User(
                id=uuid.uuid4(),
                organization_id=organization_id,
                first_name=first_name,
                last_name=last_name,
                email=email,
                password_hash=hash_password(password),
                phone=phone,
                status=UserStatus.ACTIVE,
            )
        )
        UserRoleRepository(db).assign(user.id, role.id)


def process_import_job(job_id: uuid.UUID, db: Session | None = None) -> None:
    """The row-by-row worker logic. In production the thin Celery task
    (app/tasks/imports.py) calls this with no session, so it opens and
    manages its own — this function doesn't run inside an HTTP request, so
    it can't use the request-scoped get_db dependency. Tests call it
    directly with db=<the test's session> so everything stays in the same
    in-memory database, without needing a live Celery worker.
    """
    owns_session = db is None
    session = db if db is not None else SessionLocal()
    try:
        _run(session, job_id)
    finally:
        if owns_session:
            session.close()


def _run(db: Session, job_id: uuid.UUID) -> None:
    job_repo = ImportJobRepository(db)
    error_repo = ImportRowErrorRepository(db)

    job = job_repo.get_by_id(job_id)
    if job is None:
        logger.error("process_import_job: job %s not found", job_id)
        return
    if job.status not in (ImportStatus.QUEUED, ImportStatus.PROCESSING):
        logger.warning("process_import_job: job %s has status %s, skipping", job_id, job.status)
        return

    job.status = ImportStatus.PROCESSING
    job_repo.add(job)
    db.commit()

    try:
        content = get_storage_adapter().download(job.file_url)
    except Exception as exc:  # noqa: BLE001
        job.status = ImportStatus.FAILED
        job.error_summary = f"Could not read the uploaded file: {exc}"
        job.completed_at = datetime.now(UTC)
        job_repo.add(job)
        db.commit()
        return

    mapping = job.column_mapping or {}
    seen_skus: set[str] = set()
    seen_keys: set[tuple[str, str]] = set()
    seen_emails: set[str] = set()
    success_rows = 0
    failed_rows = 0
    row_number = 1  # header is row 1; the first data row is row 2

    try:
        for row in iter_rows(content, job.file_type):
            row_number += 1
            try:
                if job.entity_type == ImportEntityType.PRODUCTS:
                    _process_products_row(
                        db,
                        organization_id=job.organization_id,
                        mapping=mapping,
                        row=row,
                        seen_skus=seen_skus,
                    )
                elif job.entity_type == ImportEntityType.USERS:
                    _process_users_row(
                        db,
                        organization_id=job.organization_id,
                        mapping=mapping,
                        row=row,
                        seen_emails=seen_emails,
                    )
                else:
                    _process_inventory_row(
                        db,
                        organization_id=job.organization_id,
                        job=job,
                        mapping=mapping,
                        row=row,
                        seen_keys=seen_keys,
                    )
            except _RowValidationError as exc:
                failed_rows += 1
                error_repo.add(
                    ImportRowError(
                        id=uuid.uuid4(),
                        import_job_id=job.id,
                        row_number=row_number,
                        field_name=exc.field_name,
                        error_message=exc.message,
                        raw_data=row,
                    )
                )
            except Exception as exc:  # noqa: BLE001 — isolate unexpected per-row failures
                logger.exception("import job %s: unexpected error on row %s", job.id, row_number)
                failed_rows += 1
                error_repo.add(
                    ImportRowError(
                        id=uuid.uuid4(),
                        import_job_id=job.id,
                        row_number=row_number,
                        field_name=None,
                        error_message=str(exc),
                        raw_data=row,
                    )
                )
            else:
                success_rows += 1

            if row_number % _COMMIT_EVERY == 0:
                job.success_rows = success_rows
                job.failed_rows = failed_rows
                job.total_rows = success_rows + failed_rows
                job_repo.add(job)
                db.commit()
    except Exception as exc:  # noqa: BLE001 — catastrophic failure (bad file structure mid-stream, etc.)
        db.rollback()
        job = job_repo.get_by_id(job_id)
        job.status = ImportStatus.FAILED
        job.error_summary = f"Import stopped unexpectedly: {exc}"
        job.completed_at = datetime.now(UTC)
        job_repo.add(job)
        db.commit()
        logger.exception("import job %s failed", job_id)
        return

    job.success_rows = success_rows
    job.failed_rows = failed_rows
    job.total_rows = success_rows + failed_rows
    job.status = ImportStatus.COMPLETED_WITH_ERRORS if failed_rows else ImportStatus.COMPLETED
    job.error_summary = f"{failed_rows} of {job.total_rows} rows failed." if failed_rows else None
    job.completed_at = datetime.now(UTC)
    job_repo.add(job)
    db.commit()
